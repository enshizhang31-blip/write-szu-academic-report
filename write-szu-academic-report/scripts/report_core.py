"""Scoped SZU DOCX authoring and structural checks. No semantic inference."""
from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

PREFIX = 'SZUHeading'
FORMATS = ['chineseCounting', 'chineseCounting', 'decimal', 'decimal', 'decimalEnclosedCircle']
LABELS = ['%1、', '(%2)', '%3.', '(%4)', '%5']
LAYOUT_P = {'keepNext', 'keepLines', 'widowControl', 'pageBreakBefore'}
ORDERS = {
    'pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
    'rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange',
    'tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange',
    'tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark headers cellIns cellDel cellMerge tcPrChange',
    'trPr': 'cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden ins del trPrChange',
    'sectPr': 'headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange',
}


def node(parent, name, **attrs):
    found = parent.find(qn('w:' + name))
    if found is None:
        found = OxmlElement('w:' + name)
        parent.append(found)
    for k, v in attrs.items():
        found.set(qn('w:' + k), str(v))
    order = ORDERS.get(etree.QName(parent).localname)
    if order:
        ranks = {name:i for i,name in enumerate(order.split())}
        parent[:] = sorted(parent, key=lambda el:ranks.get(etree.QName(el).localname,999))
    return found


def value(el, name, attr='val'):
    n = el.find(qn('w:' + name)) if el is not None else None
    return n.get(qn('w:' + attr)) if n is not None else None


def text(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t')))


def pid(el):
    return el.get(qn('w14:paraId'))


def paragraphs(doc):
    return [Paragraph(p, doc) for p in doc.element.body.iter(qn('w:p'))]


def tables(doc):
    return [Table(t, doc) for t in doc.element.body.iter(qn('w:tbl'))]


def table_id(t):
    return value(t._tbl.tblPr, 'tblCaption')


def ancestors(el, tag):
    return [a for a in el.iterancestors() if a.tag == qn('w:' + tag)]


def table_role(t):
    desc = value(t._tbl.tblPr, 'tblDescription')
    if desc == 'SZU_CODE_BOX' or any(value(p.find(qn('w:pPr')), 'pStyle') == 'CodeBox'
                                   for p in t._tbl.findall('./' + qn('w:tr') + '/' + qn('w:tc') + '/' + qn('w:p'))):
        return 'code'
    if desc == 'SZU_TEMPLATE_FRAME':
        return 'frame'
    return 'data'


def canon(el, omit=()):
    if el is None:
        return None
    el = deepcopy(el)
    for e in list(el.iter()):
        local = etree.QName(e).localname
        if local in omit and e is not el:
            e.getparent().remove(e)
            continue
        for a in list(e.attrib):
            if etree.QName(a).localname.startswith('rsid') or a == qn('w14:paraId'):
                del e.attrib[a]
    # Structural tuple avoids namespace declarations and serializer prefix noise.
    def pack(e):
        return [e.tag, [[k,v] for k,v in sorted(e.attrib.items())], e.text, [pack(c) for c in e]]
    return pack(el)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')


def load_manifest(path):
    return json.loads(Path(path).read_text(encoding='utf-8')) if path else None


def save_new(doc, source, output):
    if source and Path(source).resolve() == Path(output).resolve():
        raise ValueError('Output must differ from input; original files are never overwritten.')
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def prepare(source, output, manifest_path, layout_tables=(), fields=(), page_frame=False, heading_parents=(), upgrade_compatibility=False):
    doc = Document(source)
    settings_before=canon(doc.settings.element)
    modes=doc.settings.element.xpath('.//w:compatSetting[@w:name="compatibilityMode"]')
    legacy=bool(modes and int(modes[0].get(qn('w:val'),'15'))<14)
    if legacy and not upgrade_compatibility:
        raise ValueError('Legacy Word compatibility mode rejects table metadata. Inspect layout and explicitly pass --upgrade-compatibility to prepare a mode-15 COPY; no silent upgrade.')
    if legacy:modes[0].set(qn('w:val'),'15')
    ts = tables(doc)
    ps = paragraphs(doc)
    if any(i < 0 or i >= len(ts) for i in layout_tables):
        raise ValueError('Unknown layout-table index; inspect profile first.')
    if any(i < 0 or i >= len(ps) for i in fields):
        raise ValueError('Unknown field paragraph index; inspect outline first.')
    ids = {}
    for i, t in enumerate(ts):
        ident = 'SZU_T_' + uuid4().hex
        node(t._tbl.tblPr, 'tblCaption', val=ident)
        ids[t._tbl] = ident
    records = []
    for i, p in enumerate(ps):
        ident = f'{(uuid4().int & 0x7fffffff) or 1:08X}'
        p._p.set(qn('w14:paraId'), ident)
        parent_tables = [ids[x] for x in ancestors(p._p, 'tbl')]
        layout = any(ts[j]._tbl in ancestors(p._p, 'tbl') for j in layout_tables)
        records.append({'id': ident, 'index': i, 'text': p.text,
                        'role': 'field' if i in fields else 'protected',
                        'layout': layout, 'tables': parent_tables,
                        'xml': canon(p._p, LAYOUT_P if layout else ()),
                        'ppr': canon(p._p.pPr, LAYOUT_P if layout else ()),
                        'runs': [canon(r._r) for r in p.runs]})
    tr = []
    for i, t in enumerate(ts):
        tr.append({'id': table_id(t), 'layout': i in layout_tables,
                   'properties': canon(t._tbl.tblPr, ('tblCaption', 'tblBorders') if i in layout_tables else ('tblCaption',)),
                   'grid': canon(t._tbl.tblGrid),
                   'rows': [[canon(r._tr.trPr), [canon(c._tc.tcPr) for c in r.cells]] for r in t.rows],
                   'ancestor_tables': [ids[a] for a in ancestors(t._tbl, 'tbl')]})
    parts = {}; relationships = {}
    with ZipFile(source) as z:
        for n in z.namelist():
            if re.match(r'word/(header\d+|footer\d+|footnotes|endnotes)\.xml$', n) or n.startswith('word/media/'):
                parts[n] = hashlib.sha256(z.read(n)).hexdigest()
        original_numbering = canon(etree.fromstring(z.read('word/numbering.xml'))) if 'word/numbering.xml' in z.namelist() else None
        for n in z.namelist():
            if n.endswith('.rels'):
                relationships[n] = [dict(x.attrib) for x in etree.fromstring(z.read(n))]
    context = {}
    for pair in heading_parents:
        idx, level = map(int, pair.split(':'))
        if not 0 <= idx < len(ps) or not 1 <= level <= 5:
            raise ValueError('parent-heading must be paragraph-index:level (0-based index; level 1..5).')
        context[pid(ps[idx]._p)] = level
    manifest = {'version': 1, 'template': str(Path(source).resolve()), 'template_sha256': digest(source),
                'heading_context': context,
                'paragraphs': records, 'tables': tr, 'allow_page_frame': page_frame,
                'sections': [canon(s._sectPr, ('pgBorders',) if page_frame else ()) for s in doc.sections],
                'styles': {s.style_id: canon(s.element) for s in doc.styles},
                'doc_defaults': canon(doc.styles.element.find(qn('w:docDefaults'))),
                'relationships': relationships,
                'numbering': original_numbering, 'parts': parts,
                'settings':canon(doc.settings.element),
                'compatibility_transition':{'from_settings':settings_before,'to_mode':15,'explicit':True} if legacy else None}
    save_new(doc, source, output)
    manifest['prepared_sha256'] = digest(output)
    write_json(manifest_path, manifest)
    return {'output': str(output), 'manifest': str(manifest_path), 'protected_paragraphs': len(records)}


def issues_result(issues, **extras):
    return {**extras, 'errors': sum(i['severity'] == 'error' for i in issues),
            'warnings': sum(i['severity'] == 'warning' for i in issues), 'issues': issues}


def issue(items, rule, severity='error', **kw):
    items.append({'severity': severity, 'rule': rule, **kw})


def audit_protection(doc, manifest, path=None):
    items = []
    ps = paragraphs(doc)
    seen = [pid(p._p) for p in ps if pid(p._p)]
    if len(seen) != len(set(seen)):
        issue(items, 'duplicate_paragraph_id')
    by_id = {pid(p._p): p for p in ps}
    expected_order = [r['id'] for r in manifest['paragraphs']]
    observed_order = [i for i in seen if i in set(expected_order)]
    if observed_order != expected_order:
        issue(items, 'protected_paragraph_missing_or_reordered')
    for r in manifest['paragraphs']:
        p = by_id.get(r['id'])
        if p is None:
            continue
        ancestry = [value(a.find(qn('w:tblPr')), 'tblCaption') for a in ancestors(p._p, 'tbl')]
        if ancestry != r['tables']:
            issue(items, 'protected_paragraph_container_changed', id=r['id'])
        omit = LAYOUT_P if r['layout'] else ()
        if r['role'] == 'field':
            current = [canon(run._r) for run in p.runs]
            pos = 0
            for old in r['runs']:
                found = next((j for j in range(pos, len(current)) if current[j] == old), None)
                if found is None:
                    issue(items, 'field_original_label_or_format_changed', id=r['id'])
                    break
                pos = found + 1
            if canon(p._p.pPr, omit) != r['ppr']:
                issue(items, 'field_paragraph_format_changed', id=r['id'])
        elif canon(p._p, omit) != r['xml']:
            issue(items, 'protected_paragraph_changed', id=r['id'], text=r['text'][:60])
    ts = tables(doc)
    current_tables = {table_id(t): t for t in ts}
    active_records=[r for r in manifest['tables'] if not r.get('retired_by_reflow')]
    order = [table_id(t) for t in ts if table_id(t) in {r['id'] for r in active_records}]
    if order != [r['id'] for r in active_records]:
        issue(items, 'protected_table_missing_or_reordered')
    for r in manifest['tables']:
        t = current_tables.get(r['id'])
        if r.get('retired_by_reflow'):
            if t is not None:issue(items,'retired_layout_table_reintroduced',table=r['id'])
            continue
        if t is None:
            continue
        omit = ('tblCaption', 'tblBorders') if r['layout'] else ('tblCaption',)
        if canon(t._tbl.tblPr, omit) != r['properties'] or canon(t._tbl.tblGrid) != r['grid']:
            issue(items, 'protected_table_properties_changed', table=r['id'])
        if [value(a.find(qn('w:tblPr')), 'tblCaption') for a in ancestors(t._tbl, 'tbl')] != r['ancestor_tables']:
            issue(items, 'protected_table_container_changed', table=r['id'])
        if not r['layout'] and [[canon(row._tr.trPr), [canon(c._tc.tcPr) for c in row.cells]] for row in t.rows] != r['rows']:
            issue(items, 'protected_table_structure_changed', table=r['id'])
    sections = [canon(s._sectPr, ('pgBorders',) if manifest['allow_page_frame'] else ()) for s in doc.sections]
    if sections != manifest['sections']:
        issue(items, 'protected_section_changed')
    if 'settings' in manifest and canon(doc.settings.element)!=manifest['settings']:
        issue(items,'protected_settings_changed')
    styles = {s.style_id: canon(s.element) for s in doc.styles}
    if canon(doc.styles.element.find(qn('w:docDefaults'))) != manifest.get('doc_defaults'):
        issue(items, 'protected_document_defaults_changed')
    for sid, original in manifest['styles'].items():
        if styles.get(sid) != original:
            issue(items, 'protected_style_changed', style=sid)
    if manifest.get('numbering'):
        current = canon(doc.part.numbering_part.element)
        for original in manifest['numbering'][3]:
            if original not in current[3]:
                issue(items, 'protected_numbering_changed')
                break
    if path:
        with ZipFile(path) as z:
            for name, original_rels in manifest.get('relationships',{}).items():
                current_rels = [dict(x.attrib) for x in etree.fromstring(z.read(name))] if name in z.namelist() else []
                if any(r not in current_rels for r in original_rels):
                    issue(items, 'protected_relationship_changed', part=name)
            for name, sha in manifest['parts'].items():
                if name not in z.namelist() or hashlib.sha256(z.read(name)).hexdigest() != sha:
                    issue(items, 'protected_package_part_changed', part=name)
    # New content must reside inside an explicitly allowed template region.
    allowed = {r['id'] for r in manifest['tables'] if r['layout'] and not r.get('retired_by_reflow')}
    from report_reflow import flow_members
    flow,flow_issues=flow_members(doc,manifest);items.extend(flow_issues)
    original = set(expected_order)
    for p in ps:
        if pid(p._p) in original or not (p.text.strip() or p._p.xpath('.//w:drawing | .//w:pict')):
            continue
        parent_ids = {value(a.find(qn('w:tblPr')), 'tblCaption') for a in ancestors(p._p, 'tbl')}
        if not parent_ids.intersection(allowed) and p._p not in flow:
            issue(items, 'generated_content_outside_allowed_region', text=p.text[:60])
        for cell in ancestors(p._p, 'tc'):
            direct_original = [r for r in manifest['paragraphs'] if r['id'] in {pid(x) for x in cell.findall(qn('w:p'))}]
            if any(re.match(r'^\s*(指导教师批阅意见|成绩评定|指导教师签字|备注)', r['text']) for r in direct_original):
                issue(items, 'generated_content_in_teacher_region', text=p.text[:60])
                break
    return items


def set_font(run_or_style, size, family='SimSun', bold=False, latin=None):
    font = run_or_style.font
    font.name = latin or family
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = None
    rpr = run_or_style._element.get_or_add_rPr()
    rf = node(rpr, 'rFonts', ascii=latin or family, hAnsi=latin or family, eastAsia=family, cs=latin or family)
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        rf.attrib.pop(qn('w:' + attr), None)
    node(rpr, 'color', val='auto').attrib.pop(qn('w:themeColor'), None)
    node(rpr, 'szCs', val=int(size * 2))
    node(rpr, 'bCs', val='1' if bold else '0')


def ensure_styles(doc):
    for level in range(1, 6):
        name = PREFIX + str(level)
        st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = None
        set_font(st, 14 if level == 1 else 12, 'SimHei' if level == 1 else 'SimSun', level == 1)
        ppr = st.element.get_or_add_pPr()
        node(ppr, 'outlineLvl', val=level-1)
        node(ppr, 'ind', left=480 if level >= 3 else 0, leftChars=200 if level >= 3 else 0, firstLine=0, firstLineChars=0)
        for key in ('hanging', 'hangingChars'):
            ppr.find(qn('w:ind')).attrib.pop(qn('w:' + key), None)
        node(ppr, 'keepNext')
        node(ppr, 'keepLines')
        node(ppr, 'spacing', before=120, after=60, line=240, lineRule='auto')
    for name, size in [('SZUBody', 10.5), ('SZUCaption', 10.5), ('SZUTable', 10.5), ('CodeBox', 9)]:
        st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = None
        set_font(st, size, latin='Times New Roman')
        pf = st.paragraph_format
        pf.first_line_indent = Pt(21 if name == 'SZUBody' else 0)
        pf.left_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(6 if name == 'SZUBody' else 0)
        pf.line_spacing = 1.25 if name == 'SZUBody' else 1.0
        if name == 'SZUCaption':
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        node(st.element.get_or_add_pPr(), 'outlineLvl', val=9)


def new_numbering(doc, protected_styles=()):
    root = doc.part.numbering_part.element
    abstract_id = max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    num_id = max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))] + [0]) + 1
    ab = OxmlElement('w:abstractNum'); ab.set(qn('w:abstractNumId'), str(abstract_id))
    node(ab, 'multiLevelType', val='multilevel')
    for i in range(5):
        lv = OxmlElement('w:lvl'); lv.set(qn('w:ilvl'), str(i)); ab.append(lv)
        node(lv, 'start', val=1); node(lv, 'numFmt', val=FORMATS[i])
        # Omission uses OOXML's default: restart after the immediately preceding
        # level. Explicit lvlRestart produced missing labels in desktop Word.
        node(lv, 'pStyle', val=PREFIX + str(i+1))
        node(lv, 'suff', val='nothing'); node(lv, 'lvlText', val=LABELS[i]); node(lv, 'lvlJc', val='left')
        ppr = node(lv, 'pPr'); node(ppr, 'ind', left=480 if i >= 2 else 0, leftChars=200 if i >= 2 else 0, firstLine=0, firstLineChars=0)
        rpr = node(lv, 'rPr'); family = 'SimHei' if i == 0 else 'SimSun'
        node(rpr, 'rFonts', ascii=family, hAnsi=family, eastAsia=family)
        node(rpr, 'b', val=1 if i == 0 else 0); node(rpr, 'sz', val=28 if i == 0 else 24)
    first_num = root.find(qn('w:num'))
    if first_num is not None: first_num.addprevious(ab)
    else: root.append(ab)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id)); node(num, 'abstractNumId', val=abstract_id); root.append(num)
    for i in range(5):
        if PREFIX+str(i+1) in protected_styles:continue
        pr = doc.styles[PREFIX+str(i+1)].element.get_or_add_pPr()
        np = node(pr,'numPr'); node(np,'ilvl',val=i); node(np,'numId',val=num_id)
    return num_id


def apply_heading(p, level, num_id):
    if not 1 <= level <= 5:
        raise ValueError('Heading level must be 1..5')
    p.style = PREFIX + str(level)
    ppr = p._p.get_or_add_pPr()
    node(ppr, 'outlineLvl', val=level-1)
    np = node(ppr, 'numPr'); node(np, 'ilvl', val=level-1); node(np, 'numId', val=num_id)
    ind = node(ppr, 'ind', left=480 if level >= 3 else 0, leftChars=200 if level >= 3 else 0, firstLine=0, firstLineChars=0)
    for a in ('hanging', 'hangingChars'):
        ind.attrib.pop(qn('w:' + a), None)
    node(ppr, 'keepNext', val=1); node(ppr, 'keepLines', val=1)
    node(ppr, 'pageBreakBefore', val=0)
    node(ppr, 'spacing', before=120, after=60, line=240, lineRule='auto')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        set_font(r, 14 if level == 1 else 12, 'SimHei' if level == 1 else 'SimSun', level == 1)


def heading_level(p):
    sid = p.style.style_id
    m = re.fullmatch(PREFIX + r'([1-5])', sid)
    if m:
        return int(m[1])
    lvl = effective_p(p, 'outlineLvl', 'val')
    if lvl is not None and lvl.isdigit() and 0 <= int(lvl) <= 4:
        return int(lvl) + 1
    return None


def style_chain(style):
    seen = set(); result = []
    while style and style.style_id not in seen:
        seen.add(style.style_id); result.append(style); style = style.base_style
    return list(reversed(result))


def effective_p(p, tag, attr):
    result = None
    for st in style_chain(p.style):
        v = value(st.element.pPr, tag, attr)
        if v is not None:
            result = v
    v = value(p._p.pPr, tag, attr)
    return v if v is not None else result


def effective_run(doc, p, r):
    nodes = []
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        rd = defaults.find(qn('w:rPrDefault'))
        if rd is not None: nodes.append(rd.find(qn('w:rPr')))
    nodes += [st.element.rPr for st in style_chain(p.style)]
    if r.style: nodes += [st.element.rPr for st in style_chain(r.style)]
    nodes.append(r._r.rPr)
    result = {}
    for pr in nodes:
        if pr is None: continue
        for key, tag, attr in [('size', 'sz', 'val'), ('bold', 'b', 'val'), ('eastAsia','rFonts','eastAsia'), ('ascii','rFonts','ascii'), ('hAnsi','rFonts','hAnsi'), ('color','color','val')]:
            el = pr.find(qn('w:' + tag))
            if el is not None:
                v = el.get(qn('w:' + attr))
                if tag == 'b' and v is None: v = '1'
                if v is not None: result[key] = v
        rf = pr.find(qn('w:rFonts'))
        if rf is not None and all(rf.get(qn('w:'+a)) for a in ('ascii','hAnsi','eastAsia')):
            result.pop('theme_font', None)
        if rf is not None and any(rf.get(qn('w:'+a)) for a in ('asciiTheme','hAnsiTheme','eastAsiaTheme')):
            result['theme_font'] = True
        col = pr.find(qn('w:color'))
        if col is not None and col.get(qn('w:val')) and not col.get(qn('w:themeColor')):
            result.pop('theme_color', None)
        if col is not None and col.get(qn('w:themeColor')):
            result['theme_color'] = col.get(qn('w:themeColor'))
    return result


def outline(doc, manifest=None):
    protected = {r['id'] for r in manifest['paragraphs']} if manifest else set()
    entries = []
    for i, p in enumerate(paragraphs(doc)):
        if p.text.strip() or p._p.xpath('.//w:drawing | .//w:pict'):
            is_protected = pid(p._p) in protected
            context = manifest.get('heading_context',{}) if manifest else {}
            entries.append({'index': i, 'id': pid(p._p), 'text': p.text,
                            'observed_level': heading_level(p),
                            'level': context.get(pid(p._p)) if is_protected else heading_level(p), 'protected': is_protected,
                            'question': '', 'evidence': [], 'parent_index': None})
    return {'notice': 'Existing levels are diagnostic only. Rebuild from content; explicitly set parent_index for each subheading.',
            'paragraphs': entries, 'semantic_review': 'pending'}


def validate_outline(entries):
    items = []; stack = {}
    for e in entries:
        level = e.get('level')
        if level is None: continue
        if not isinstance(level, int) or isinstance(level, bool) or not 1 <= level <= 5:
            issue(items, 'invalid_heading_level', index=e.get('index')); continue
        if level > 1 and level-1 not in stack:
            issue(items, 'heading_parent_missing', index=e.get('index'))
        expected = stack.get(level-1) if level > 1 else None
        if e.get('parent_index') != expected:
            issue(items, 'heading_parent_mismatch', index=e.get('index'), expected=expected)
        if not e.get('question', '').strip():
            issue(items, 'heading_question_missing', index=e.get('index'))
        stack = {k:v for k,v in stack.items() if k < level}; stack[level] = e['index']
    return items


def set_borders(t, outer=True):
    b = node(t._tbl.tblPr, 'tblBorders')
    for old in list(b):b.remove(old)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node(b, edge, val='single' if outer or edge.startswith('inside') else 'nil', sz=4, space=0, color='000000')


def format_table(t, role):
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = None
    # WPS treats an explicit zero cell-spacing element as separate border mode.
    # Omit it to use the collapsed grid in both editors.
    spacing=t._tbl.tblPr.find(qn('w:tblCellSpacing'))
    if spacing is not None:t._tbl.tblPr.remove(spacing)
    if value(t._tbl.tblPr,'tblLayout','type')=='fixed':
        widths=[int(c.get(qn('w:w'))) for c in t._tbl.tblGrid if c.get(qn('w:w'))]
        if widths:node(t._tbl.tblPr,'tblW',w=sum(widths),type='dxa')
    set_borders(t)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        spacing=pr.find(qn('w:tblCellSpacing'))
        if spacing is not None:pr.remove(spacing)
        height = pr.find(qn('w:trHeight'))
        if height is not None: pr.remove(height)
        if role == 'code' or sum(len(c.text) for c in row.cells)<900:
            node(pr, 'cantSplit')
        else:
            split = pr.find(qn('w:cantSplit'))
            if split is not None: pr.remove(split)
        if i == 0 and role == 'data': node(pr, 'tblHeader')
        for cell in row.cells:
            cp = cell._tc.get_or_add_tcPr()
            border = cp.find(qn('w:tcBorders'))
            if border is not None: cp.remove(border)
            # One source of grid lines prevents duplicated strokes in WPS.
            # Use table-level borders; no redundant cell-level copies.
            node(cp, 'shd', fill='F7F7F7' if role == 'code' else ('F2F2F2' if i > 0 and i % 2 == 0 else 'FFFFFF'))
            margins = node(cp, 'tcMar')
            for side in ('top','left','bottom','right'): node(margins, side, w=100, type='dxa')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def page_frame(doc, cover=True):
    for index,section in enumerate(doc.sections):
        b = node(section._sectPr, 'pgBorders', offsetFrom='text', display='notFirstPage' if cover and index==0 else 'allPages')
        for side in ('top','left','bottom','right'):
            node(b, side, val='single', sz=4, space=4, color='000000')


def format_document(source, output, manifest=None, outline_data=None, generated_all=False, frame=False):
    if manifest is None and not generated_all:
        raise ValueError('Use --manifest for templates, or --generated-all ONLY for an entirely generated document.')
    doc = Document(source)
    if manifest:
        before = audit_protection(doc, manifest, source)
        if before: raise ValueError('Template protection failed before formatting: ' + json.dumps(before[:3], ensure_ascii=False))
    reserved = {PREFIX+str(i) for i in range(1,6)} | {'SZUBody','SZUCaption','SZUTable','CodeBox'}
    if manifest and reserved.intersection(manifest['styles']):
        # Avoid changing inherited template formatting through shared styles.
        probe = Document(source); ensure_styles(probe)
        for name in reserved.intersection(manifest['styles']):
            if canon(probe.styles[name].element) != manifest['styles'][name]:
                raise ValueError('Template owns incompatible reserved style: ' + name)
    ps = paragraphs(doc)
    mapped = {}
    if outline_data:
        entries = outline_data['paragraphs']
        problems = validate_outline(entries)
        if problems: raise ValueError('Invalid outline: ' + json.dumps(problems, ensure_ascii=False))
        for e in entries:
            i = e['index']
            if not isinstance(i, int) or i < 0 or i >= len(ps) or ps[i].text != e['text']:
                raise ValueError('Stale outline; extract again after editing paragraphs.')
            if i in mapped: raise ValueError('Duplicate outline paragraph index')
            mapped[i] = e
    protected = {r['id'] for r in manifest['paragraphs']} if manifest else set()
    for i,e in mapped.items():
        if pid(ps[i]._p) in protected and e.get('level') != manifest.get('heading_context',{}).get(pid(ps[i]._p)):
            raise ValueError('Protected heading level must match prepare --parent-heading context; it will not be restyled.')
    ensure_styles(doc); num_id = new_numbering(doc, manifest['styles'] if manifest else ())
    for i, p in enumerate(ps):
        if pid(p._p) in protected: continue
        # Only explicit plan levels or existing semantic styles; never infer from numeric text.
        level = mapped[i].get('level') if i in mapped else heading_level(p)
        if level:
            if i in mapped and 'replacement_text' in mapped[i]:
                replacement = mapped[i]['replacement_text']
                if not isinstance(replacement, str) or not replacement.strip():
                    raise ValueError('replacement_text must be nonempty heading text.')
                p.text = replacement
            if re.match(r'^\s*(?:[一二三四五六七八九十]+、|\([一二三四五六七八九十\d]+\)|\d+\.|[①-⑳])', p.text):
                raise ValueError('Remove manually typed numbering from heading before applying automatic numbering: '+p.text[:60])
            apply_heading(p, level, num_id); continue
        tbls = ancestors(p._p, 'tbl')
        nearest = Table(tbls[0], doc) if tbls else None
        role = table_role(nearest) if nearest is not None else None
        if nearest is not None and manifest and table_id(nearest) in {r['id'] for r in manifest['tables'] if r['layout']}:
            role = 'frame'
        visual = bool(p._p.xpath('.//w:drawing | .//w:pict'))
        caption = p.style.name == 'SZUCaption'
        code = role == 'code' or p.style.name == 'CodeBox'
        style = 'CodeBox' if code else ('SZUCaption' if caption else ('SZUTable' if role == 'data' else 'SZUBody'))
        p.style = style
        ppr = p._p.get_or_add_pPr()
        for tag in ('numPr', 'outlineLvl'):
            n = ppr.find(qn('w:'+tag))
            if n is not None: ppr.remove(n)
        pf = p.paragraph_format
        pf.left_indent = Pt(0); pf.right_indent = Pt(0)
        pf.first_line_indent = Pt(21 if style == 'SZUBody' and not visual else 0)
        ind = node(ppr, 'ind')
        for attr in ('leftChars','rightChars','firstLineChars','hangingChars','hanging'):
            ind.attrib.pop(qn('w:'+attr), None)
        pf.keep_with_next = visual
        pf.keep_together = caption or visual
        pf.widow_control = True
        pf.space_before = Pt(0); pf.space_after = Pt(6 if style == 'SZUBody' and not visual else 0)
        pf.line_spacing = 1.25 if style == 'SZUBody' and not visual else 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if caption or visual else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs: set_font(run, 9 if code else 10.5, latin='Times New Roman')
    original_tables = {r['id']: r for r in manifest['tables']} if manifest else {}
    for t in tables(doc):
        r = original_tables.get(table_id(t))
        if r:
            if r['layout']:
                # Layout wrappers are diagnosed/reflowed, never restyled as data grids.
                for row in t.rows:
                    pr = row._tr.get_or_add_trPr()
                    cs = pr.find(qn('w:cantSplit'))
                    if cs is not None: pr.remove(cs)
                    # Outer cells must not suppress the frame's explicit edges.
                    for c in row.cells:
                        cb = c._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
                        if cb is not None: c._tc.tcPr.remove(cb)
            continue
        format_table(t, table_role(t))
    if frame:
        if manifest and not manifest['allow_page_frame']:
            raise ValueError('Page frame was not allowed by prepare --page-frame.')
        page_frame(doc)
    save_new(doc, source, output)
    if manifest:
        after = audit_protection(Document(output), manifest, output)
        if after: raise ValueError('Formatted output failed template protection: '+json.dumps(after[:3], ensure_ascii=False))
    return {'output': str(output), 'semantic_review': 'pending', 'visual_review': 'pending'}


def audit(doc, manifest=None, path=None):
    items = audit_protection(doc, manifest, path) if manifest else []
    if manifest is None:
        issue(items, 'scope_assumed_entire_document_generated', 'warning')
    protected = {r['id'] for r in manifest['paragraphs']} if manifest else set()
    protected_tables = {r['id'] for r in manifest['tables']} if manifest else set()
    layout_tables = {r['id'] for r in manifest['tables'] if r['layout']} if manifest else set()
    stack = {}; counts = {'headings':0, 'paragraphs':0, 'code_boxes':0}
    numbering = doc.part.numbering_part.element
    used_num = None
    for i, p in enumerate(paragraphs(doc)):
        if pid(p._p) in protected:
            context_level = manifest.get('heading_context',{}).get(pid(p._p))
            if context_level:
                stack = {k:v for k,v in stack.items() if k < context_level}
                stack[context_level] = i
            continue
        if not p.text.strip() and not p._p.xpath('.//w:drawing | .//w:pict'): continue
        counts['paragraphs'] += 1
        level = heading_level(p)
        parent_tables = ancestors(p._p, 'tbl')
        t = Table(parent_tables[0], doc) if parent_tables else None
        role = table_role(t) if t is not None else 'body'
        if t is not None and table_id(t) in layout_tables: role = 'body'
        visual = bool(p._p.xpath('.//w:drawing | .//w:pict'))
        caption = p.style.name == 'SZUCaption'
        if caption or visual:
            if effective_p(p, 'jc', 'val') != 'center': issue(items, 'figure_caption_not_centered', paragraph=i)
        if caption and re.match(r'^[图表]\s+\d', p.text): issue(items, 'caption_label_space', paragraph=i)
        if visual and effective_p(p, 'keepNext','val') not in ('1','true', None): issue(items, 'figure_caption_keep_next_missing', paragraph=i)
        if level:
            counts['headings'] += 1
            if level > 1 and level-1 not in stack: issue(items, 'heading_parent_missing', paragraph=i)
            stack = {k:v for k,v in stack.items() if k < level}; stack[level] = i
            expected_style = PREFIX + str(level)
            if p.style.style_id != expected_style: issue(items, 'heading_style_mismatch', paragraph=i, expected=expected_style)
            if effective_p(p, 'outlineLvl','val') != str(level-1): issue(items, 'heading_outline_level', paragraph=i)
            expected_left = '480' if level >= 3 else '0'
            if effective_p(p,'ind','left') != expected_left or effective_p(p,'ind','firstLine') != '0': issue(items, 'heading_indent', paragraph=i)
            if effective_p(p,'ind','leftChars') != ('200' if level >= 3 else '0') or effective_p(p,'ind','firstLineChars') not in (None,'0') or effective_p(p,'ind','hanging') not in (None,'0'):
                issue(items, 'heading_character_indent', paragraph=i)
            np = p._p.pPr.find(qn('w:numPr')) if p._p.pPr is not None else None
            nid = value(np, 'numId'); ilvl = value(np, 'ilvl')
            if ilvl != str(level-1) or not nid: issue(items, 'heading_numbering_missing_or_wrong_level', paragraph=i)
            nums = numbering.xpath('./w:num[@w:numId="'+str(nid)+'"]')
            aid = value(nums[0], 'abstractNumId') if nums else None
            levels = numbering.xpath('./w:abstractNum[@w:abstractNumId="'+str(aid)+'"]/w:lvl[@w:ilvl="'+str(level-1)+'"]')
            if not levels or value(levels[0],'numFmt') != FORMATS[level-1] or value(levels[0],'lvlText') != LABELS[level-1] or (level > 1 and value(levels[0],'lvlRestart') is not None):
                issue(items, 'heading_numbering_definition', paragraph=i)
            if nums and nums[0].find(qn('w:lvlOverride')) is not None: issue(items, 'unexpected_numbering_override', paragraph=i)
            if used_num and nid != used_num: issue(items, 'heading_numbering_sequence_changed', paragraph=i)
            used_num = nid
            if re.match(r'^\s*(?:[一二三四五六七八九十]+、|\([一二三四五六七八九十\d]+\)|\d+\.|[①-⑳])', p.text): issue(items, 'manual_heading_numbering', paragraph=i)
            expected_size = 28 if level == 1 else 24
            expected_font = 'SimHei' if level == 1 else 'SimSun'
        else:
            expected_size = 18 if role == 'code' or p.style.name == 'CodeBox' else 21
            expected_font = 'SimSun'
            if role in ('body','frame') and not visual and not caption and p.text.strip():
                if effective_p(p,'ind','firstLine') not in ('420','421'): issue(items, 'body_two_character_indent', paragraph=i)
                if re.match(r'^\s*(?:class |def |import |public |const |function )', p.text): issue(items, 'displayed_code_outside_box', paragraph=i)
        for r in p.runs:
            if not r.text.strip(): continue
            props = effective_run(doc,p,r)
            if props.get('size') != str(expected_size): issue(items, 'font_size', paragraph=i, actual=props.get('size'), expected=expected_size/2)
            aliases = {'SimSun':{'SimSun','宋体'},'SimHei':{'SimHei','黑体'}}
            if re.search(r'[\u3400-\u9fff]',r.text) and props.get('eastAsia') not in aliases[expected_font]: issue(items, 'chinese_font', paragraph=i)
            lf = expected_font if level else 'Times New Roman'
            if re.search(r'[A-Za-z0-9]',r.text) and (props.get('ascii') not in aliases.get(lf,{lf}) or props.get('hAnsi') not in aliases.get(lf,{lf})): issue(items, 'latin_font', paragraph=i)
            if props.get('color') not in (None,'auto','000000') or props.get('theme_color'): issue(items, 'nonblack_text', paragraph=i)
            if props.get('theme_font'): issue(items, 'unresolved_theme_font', paragraph=i)
            if level and ((props.get('bold') in ('1','true')) != (level == 1)): issue(items, 'heading_bold', paragraph=i)
    for i,t in enumerate(tables(doc)):
        from report_reflow import candidate_reason
        record=next((r for r in (manifest or {}).get('tables',[]) if r['id']==table_id(t)),{})
        reason=None if record.get('role')=='form' else candidate_reason(t)
        if reason:issue(items,'body_layout_table','warning',table=i,detail=reason)
        if table_id(t) in protected_tables: continue
        role = table_role(t)
        if role == 'code': counts['code_boxes'] += 1
        b = t._tbl.tblPr.find(qn('w:tblBorders'))
        for edge in ('top','left','bottom','right') + (() if role == 'code' else ('insideH','insideV')):
            if value(b,edge)!='single':issue(items,'table_border_missing',table=i,edge=edge)
            if value(b,edge,'sz')!='4' or value(b,edge,'color')!='000000':issue(items,'table_border_spec',table=i,edge=edge,expected='0.5pt black single')
        if value(t._tbl.tblPr,'tblCellSpacing','w') not in (None,'0'):issue(items,'table_cell_spacing',table=i)
        if t._tbl.tblPr.find(qn('w:tblStyle')) is not None:issue(items,'table_style_border_conflict','warning',table=i)
        for ri,row in enumerate(t.rows):
            h = row._tr.get_or_add_trPr().find(qn('w:trHeight'))
            if h is not None and h.get(qn('w:hRule')) == 'exact': issue(items,'fixed_row_height',table=i)
            for c in row.cells:
                cb = c._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
                if cb is not None and any(x.get(qn('w:val')) in ('nil','none') for x in cb): issue(items,'cell_border_override',table=i)
                if cb is not None and any(x.get(qn('w:val'))!='single' or x.get(qn('w:sz'))!='4' or x.get(qn('w:color'))!='000000' for x in cb):issue(items,'cell_border_spec',table=i)
                if role == 'data':
                    fill = value(c._tc.tcPr,'shd','fill'); expected = 'F2F2F2' if ri > 0 and ri % 2 == 0 else 'FFFFFF'
                    if fill != expected: issue(items,'table_fill',table=i)
        if role=='data' and len(t.rows)>1 and t.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:issue(items,'table_repeat_header_missing',table=i)
    return issues_result(items, counts=counts, semantic_review='pending', visual_review='pending')
