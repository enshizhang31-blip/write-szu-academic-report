#!/usr/bin/env python3
"""Unified, local SZU report tools. Run any subcommand with --help."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from report_core import (audit, digest, format_document, heading_level, issues_result,
                         issue, load_manifest, outline, paragraphs, prepare, tables,
                         table_role, validate_outline, write_json, effective_run,
                         effective_p)
except ImportError as exc:
    print(json.dumps({'status':'blocked','error':'Missing dependency: '+str(exc),'type':'ImportError'}))
    raise SystemExit(2)


def profile(path):
    doc = Document(path)
    return {'file':str(path), 'sha256':digest(path),
            'sections':[{'width_pt':s.page_width.pt,'height_pt':s.page_height.pt,
                         'margins_pt':{k:getattr(s,k).pt for k in ('top_margin','bottom_margin','left_margin','right_margin')}} for s in doc.sections],
            'paragraphs':[{'index':i,'text':p.text,'observed_level':heading_level(p),
                           'style':p.style.name,'left':effective_p(p,'ind','left'),
                           'runs':[{'text':r.text,'effective':effective_run(doc,p,r)} for r in p.runs if r.text.strip()]}
                          for i,p in enumerate(paragraphs(doc)) if p.text.strip()],
            'tables':[{'index':i,'role':table_role(t),'rows':len(t.rows),
                       'columns':len(t.columns),'layout_warning':candidate_reason(t),'text':t._tbl.xpath('.//w:t/text()')[:6]} for i,t in enumerate(tables(doc))],
            'notice':'Observed headings are diagnostic, never an authoritative hierarchy.'}


from report_editors import render, doctor
from report_reflow import reflow, candidate_reason
from report_verify import verify


def render_risks(images, texts, headings):
    """Heuristics only: no claim of automatic visual correctness."""
    risks=[]
    for i,im in enumerate(images,1):
        # Ignore margins and thin vertical frame lines when estimating body ink.
        gray=im.convert('L'); w,h=gray.size
        crop=gray.crop((int(w*.16),int(h*.09),int(w*.84),int(h*.91)))
        rows=[]
        for y in range(crop.height):
            count=sum(v < 150 for v in crop.crop((0,y,crop.width,y+1)).tobytes())
            if count > max(10,crop.width*.025) and count < crop.width*.8: rows.append(y)
        if i not in (1,len(images)) and rows and max(rows) < crop.height*.65:
            risks.append({'page':i,'rule':'large_bottom_whitespace','review':'Check whether an unsplittable block pushed content to the next page.'})
        if not texts[i-1].strip(): risks.append({'page':i,'rule':'empty_text_or_image_only_page'})
    for hd in headings:
        p=hd.get('page',0)
        if not 1 <= p <= len(texts): continue
        flat=re.sub(r'\s+','',texts[p-1]); title=re.sub(r'\s+','',hd['text'])
        if title and flat.endswith(title): risks.append({'page':p,'rule':'possible_orphan_heading','text':hd['text']})
    return risks


def evidence_audit(path):
    issues=[]
    if not path: return issues
    data=load_manifest(path)
    for row in data.get('requirements',[]):
        status=row.get('status')
        if status not in ('complete','incomplete','not_applicable'):
            issue(issues,'invalid_requirement_status',requirement=row.get('id')); continue
        if status == 'incomplete': issue(issues,'requirement_incomplete','warning',requirement=row.get('id'))
        evidence=row.get('evidence',[])
        if status == 'complete' and not evidence: issue(issues,'completed_requirement_without_evidence',requirement=row.get('id'))
        for e in evidence:
            f=Path(e['path']); f=f if f.is_absolute() else Path(path).parent/f
            if not f.is_file(): issue(issues,'evidence_file_missing',file=str(f))
            if e.get('kind') not in ('source','log','raw_image','result','screenshot','manual_review'):
                issue(issues,'evidence_kind_unknown',requirement=row.get('id'))
        if row.get('needs_result') and status == 'complete' and not any(e.get('kind') in ('result','screenshot','manual_review') for e in evidence):
            issue(issues,'raw_input_not_result_evidence',requirement=row.get('id'))
    return issues


def source_check(docx, source):
    if not source: return []
    import audit_code_consistency as checker
    declared=checker.source_declarations(Path(source))
    mentioned=set(checker.DOC_IDENTIFIER.findall(checker.document_text(Path(docx))))-checker.ALLOW
    # Capitalization heuristics cannot prove identifiers are incorrect.
    return [{'severity':'warning','rule':'source_identifier_review','identifier':x} for x in sorted(mentioned-declared)]


def validate_review(path, report, rendered):
    if not path: return False, ['semantic_review_pending','visual_review_pending']
    review=load_manifest(path); missing=[]
    if review.get('source_sha256') != digest(report): missing.append('review_stale_source')
    if not review.get('reviewer') or not review.get('semantic_pass') or not review.get('semantic_notes'):
        missing.append('semantic_review_pending')
    if review.get('render_pdf_sha256') != rendered.get('pdf_sha256'): missing.append('review_stale_render')
    if sorted(review.get('pages_reviewed',[])) != list(range(1,rendered.get('pages',0)+1)):
        missing.append('visual_pages_incomplete')
    if not review.get('visual_pass') or not review.get('visual_notes'): missing.append('visual_review_pending')
    required={'closed_frames','continuous_table_grid','readable_figures','no_orphan_headings','figure_caption_pairs','whitespace_reviewed'}
    if not required.issubset(set(review.get('checks_passed',[]))): missing.append('visual_checklist_incomplete')
    return not missing,missing


def rendered_numbering_issues(rendered):
    """Check Word's displayed labels, not only stored numbering definitions."""
    if rendered.get('backend') not in ('word','wps'): return []
    items=[]; counters=[0]*5
    def chinese(n):
        digits='零一二三四五六七八九'
        if n < 10:return digits[n]
        if n < 100:return (digits[n//10] if n>=20 else '')+'十'+(digits[n%10] if n%10 else '')
        return None
    for h in rendered.get('headings',[]):
        level=h['level'];counters[level-1]+=1
        for i in range(level,5):counters[i]=0
        n=counters[level-1];c=chinese(n)
        expected=(c+'、' if c else None) if level==1 else (('('+c+')' if c else None) if level==2 else (str(n)+'.' if level==3 else ('('+str(n)+')' if level==4 else (chr(0x2460+n-1) if n<=20 else None))))
        if not h.get('label') or (expected and h['label']!=expected):
            issue(items,'rendered_heading_numbering',text=h['text'],actual=h.get('label'),expected=expected,page=h['page'])
    return items


def parser():
    p=argparse.ArgumentParser(description=__doc__)
    sub=p.add_subparsers(dest='command',required=True)
    for name in ('profile','outline','audit','format','render','verify'):
        q=sub.add_parser(name)
        q.add_argument('docx',type=Path)
        if name in ('outline','audit','format','verify'): q.add_argument('--manifest',type=Path)
        if name in ('profile','outline','audit','verify','render'): q.add_argument('--json',type=Path)
        if name in ('outline','format'): q.add_argument('--plan',type=Path,help='Explicit content-derived outline JSON')
        if name in ('audit','verify'):
            q.add_argument('--source-dir',type=Path); q.add_argument('--evidence',type=Path)
        if name == 'format':
            q.add_argument('--output',required=True,type=Path);q.add_argument('--generated-all',action='store_true');q.add_argument('--page-frame',action='store_true')
        if name in ('render','verify'):
            q.add_argument('--output-dir',required=True,type=Path)
            q.add_argument('--backend',choices=['auto','word','wps','libreoffice'],default='auto')
            q.add_argument('--soffice',type=Path,help='Absolute bundled LibreOffice executable')
            q.add_argument('--dpi',type=int,default=200,help='PNG resolution, 96..400; default 200')
            if name == 'render': q.add_argument('--roundtrip',action='store_true',help='Save/reopen an isolated copy and compare displayed numbering')
        if name == 'verify':
            q.add_argument('--primary-editor',choices=['wps','word','libreoffice'],default='wps')
            q.add_argument('--compat-editor',choices=['word','wps','none'],default='word')
            q.add_argument('--review',type=Path);q.add_argument('--reuse-render',action='store_true',help='Reuse render only if source and rendered file hashes match')
    q=sub.add_parser('prepare');q.add_argument('template',type=Path);q.add_argument('--output',type=Path,required=True);q.add_argument('--manifest',type=Path,required=True)
    q.add_argument('--layout-table',type=int,action='append',default=[],help='Zero-based table index; explicitly allow filling and layout changes')
    q.add_argument('--field',type=int,action='append',default=[],help='Zero-based paragraph index; original runs must remain, new runs may be added')
    q.add_argument('--parent-heading',action='append',default=[],help='Existing semantic template heading as paragraph-index:level; preserve its original formatting')
    q.add_argument('--page-frame',action='store_true',help='Allow section page borders on the working copy')
    q.add_argument('--upgrade-compatibility',action='store_true',help='Explicitly upgrade legacy mode in a COPY; requires visual recheck')
    q=sub.add_parser('doctor');q.add_argument('--json',type=Path);q.add_argument('--probe',action='store_true');q.add_argument('--output-dir',type=Path)
    q=sub.add_parser('reflow');q.add_argument('docx',type=Path);q.add_argument('--manifest',type=Path,required=True);q.add_argument('--output',type=Path,required=True);q.add_argument('--manifest-out',type=Path,required=True);q.add_argument('--layout-table',type=int,action='append',required=True);q.add_argument('--json',type=Path)
    return p


def main(argv=None):
    args=parser().parse_args(argv)
    try:
        code=0; command=args.command
        manifest=load_manifest(getattr(args,'manifest',None)) if command != 'prepare' else None
        if command == 'doctor': result=doctor(args.probe,args.output_dir);code=2 if args.probe and any(e['automation']!='export_verified' for e in result['editors']) else 0
        elif command == 'reflow': result=reflow(args.docx,args.output,manifest,args.manifest_out,args.layout_table)
        elif command == 'profile': result=profile(args.docx)
        elif command == 'prepare': result=prepare(args.template,args.output,args.manifest,args.layout_table,args.field,args.page_frame,args.parent_heading,args.upgrade_compatibility)
        elif command == 'outline':
            result=outline(Document(args.docx),manifest)
            if args.plan:
                data=load_manifest(args.plan);result['validation']=issues_result(validate_outline(data['paragraphs']))
                code=1 if result['validation']['errors'] else 0
        elif command == 'format': result=format_document(args.docx,args.output,manifest,load_manifest(args.plan),args.generated_all,args.page_frame)
        elif command == 'render': result=render(args.docx,args.output_dir,args.backend,args.soffice,args.dpi,args.roundtrip)
        elif command == 'verify':
            primary=args.primary_editor if args.backend=='auto' else args.backend
            result,code=verify(args.docx,manifest,args.output_dir,primary,args.compat_editor,args.reuse_render,args.review,args.dpi,args.soffice,args.evidence,args.source_dir)
        else:
            result=audit(Document(args.docx),manifest,args.docx)
            issues=result['issues']+evidence_audit(args.evidence)+source_check(args.docx,args.source_dir)
            result.update(issues_result(issues));code=1 if result['errors'] else 0
        if getattr(args,'json',None): write_json(args.json,result)
        # Keep large render text and profile runs in JSON files when requested.
        summary={k:v for k,v in result.items() if k not in ('page_text','render','paragraphs')}
        if getattr(args,'json',None): summary['json']=str(args.json)
        print(json.dumps(summary if getattr(args,'json',None) or command=='render' else result,ensure_ascii=False,indent=2))
        return code
    except Exception as exc:
        result={'status':'blocked','error':str(exc),'type':type(exc).__name__}
        if getattr(args,'json',None): write_json(args.json,result)
        print(json.dumps(result,ensure_ascii=False,indent=2));return 2


if __name__ == '__main__':
    sys.exit(main())
