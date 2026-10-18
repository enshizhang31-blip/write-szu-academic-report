#!/usr/bin/env python3
"""Create an explicitly unfinished starter with the five-level style system."""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Mm
from report_core import ensure_styles,new_numbering,apply_heading,set_font,page_frame

def build(output):
    doc=Document();s=doc.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297)
    s.top_margin=s.bottom_margin=Mm(25.4);s.left_margin=s.right_margin=Mm(31.7)
    ensure_styles(doc);nid=new_numbering(doc);page_frame(doc)
    doc.add_paragraph('深圳大学实验报告')
    for field in ['课程名称：','实验名称：','姓名：','学号：','学院：','实验日期：']:
        doc.add_paragraph(field)
    doc.add_page_break()
    sections=[('实验目的与要求','填写本实验的问题、目标及要求。'),('实验方法与实现','按真实实验内容建立必要的下级标题。'),('实验结果与分析','加入实际结果及对应分析，缺失证据不得编造。'),('实验结论','根据已完成实验总结结论和局限。')]
    for title,body in sections:
        p=doc.add_paragraph(title);apply_heading(p,1,nid)
        p=doc.add_paragraph(style='SZUBody');set_font(p.add_run(body),10.5,latin='Times New Roman')
    Path(output).parent.mkdir(parents=True,exist_ok=True);doc.save(output)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('output');a=p.parse_args();build(a.output);print(a.output)
