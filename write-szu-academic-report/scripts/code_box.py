#!/usr/bin/env python3
"""Insert a framed 9pt excerpt in a Document or an authorized template cell."""
from docx.shared import Pt
from report_core import ensure_styles, format_table, node, set_font

def ensure_code_box_style(doc):
    ensure_styles(doc)
    return doc.styles['CodeBox']

def add_code_box(container, code, before=None):
    doc=container.part.document
    ensure_code_box_style(doc)
    leading=container.add_paragraph()
    leading.paragraph_format.space_before=Pt(0)
    leading.paragraph_format.space_after=Pt(0)
    leading.paragraph_format.keep_with_next=True
    set_font(leading.add_run(''),10.5,latin='Times New Roman')
    t=container.add_table(rows=1,cols=1)
    node(t._tbl.tblPr,'tblDescription',val='SZU_CODE_BOX')
    format_table(t,'code')
    c=t.cell(0,0)
    for i,line in enumerate(code.rstrip('\n').splitlines() or ['']):
        p=c.paragraphs[0] if i==0 else c.add_paragraph()
        p.style='CodeBox';p.paragraph_format.first_line_indent=Pt(0)
        set_font(p.add_run(line or ' '),9,latin='Times New Roman')
    trailing=container.add_paragraph()
    trailing.paragraph_format.space_before=Pt(0)
    trailing.paragraph_format.space_after=Pt(0)
    set_font(trailing.add_run(''),10.5,latin='Times New Roman')
    if before is not None:
        for e in (leading._p,t._tbl,trailing._p):before._p.addprevious(e)
    return t
