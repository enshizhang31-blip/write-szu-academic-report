"""Behavioral regression tests; fixtures live only in a caller-selected directory."""
import argparse
import json
import tempfile
import unittest
import os
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from report_core import *
from report_tools import evidence_audit, validate_review, rendered_numbering_issues
from code_box import add_code_box


class ReportTests(unittest.TestCase):
    def setUp(self):
        base=Path(os.environ.get('SZU_TEST_ROOT',str(Path.cwd()/'work'/'szu-tests'))).resolve();base.mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=base);self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def test_all_five_levels_and_restart_definition(self):
        d=Document();ensure_styles(d);n=new_numbering(d)
        for lev in [1,2,3,4,5,4,5,3,4,5,2,3,1,2,3,4,5]:
            p=d.add_paragraph('标题内容');apply_heading(p,lev,n)
        f=self.root/'levels.docx';d.save(f);d=Document(f)
        self.assertEqual(audit(d)['errors'],0)
        self.assertEqual([heading_level(p) for p in d.paragraphs],[1,2,3,4,5,4,5,3,4,5,2,3,1,2,3,4,5])
        for p in d.paragraphs:
            level=heading_level(p)
            self.assertEqual(effective_p(p,'ind','left'),'480' if level>=3 else '0')
            self.assertEqual(effective_p(p,'outlineLvl','val'),str(level-1))
    def make_template(self):
        d=Document();d.add_paragraph('课程及实验要求：不得改写')
        d.add_paragraph('姓名：')
        t=d.add_table(rows=1,cols=1);t.cell(0,0).text='实验过程及内容：'
        d.add_paragraph('指导教师签字：')
        src=self.root/'template.docx';d.save(src)
        work=self.root/'work.docx';mf=self.root/'regions.json'
        prepare(src,work,mf,layout_tables=[0],fields=[1])
        return src,work,mf
    def test_fill_not_whole_table_equality(self):
        src,work,mf=self.make_template();m=load_manifest(mf);d=Document(work)
        d.paragraphs[1].add_run('测试学生')
        d.tables[0].cell(0,0).add_paragraph('新增正文。')
        draft=self.root/'draft.docx';d.save(draft)
        self.assertEqual(audit_protection(Document(draft),m,draft),[])
        out=self.root/'formatted.docx';format_document(draft,out,m)
        self.assertEqual(audit(Document(out),m,out)['errors'],0)
    def test_protected_changes_detected(self):
        src,work,mf=self.make_template();m=load_manifest(mf)
        for target in (0,2):
            d=Document(work);d.paragraphs[target].runs[0].bold=True
            self.assertTrue(audit_protection(d,m))
        d=Document(work);d.tables[0].cell(0,0).paragraphs[0].runs[0].text='改掉实验要求'
        self.assertTrue(audit_protection(d,m))
    def test_generated_inside_frame_is_not_exempt(self):
        src,work,mf=self.make_template();m=load_manifest(mf);d=Document(work)
        p=d.tables[0].cell(0,0).add_paragraph('框内错误字体')
        p.runs[0].font.size=Pt(20)
        self.assertIn('font_size',{i['rule'] for i in audit(d,m)['issues']})
    def test_scope_and_duplicate_identity(self):
        src,work,mf=self.make_template();m=load_manifest(mf);d=Document(work)
        p=d.add_paragraph('框外生成正文')
        self.assertIn('generated_content_outside_allowed_region',{i['rule'] for i in audit_protection(d,m)})
        p._p.set(qn('w14:paraId'),pid(d.paragraphs[0]._p))
        self.assertIn('duplicate_paragraph_id',{i['rule'] for i in audit_protection(d,m)})
    def test_code_in_cell_remains_nine_points(self):
        src,work,mf=self.make_template();m=load_manifest(mf);d=Document(work)
        add_code_box(d.tables[0].cell(0,0),'def run():\n    return 1')
        draft=self.root/'code.docx';d.save(draft);out=self.root/'formatted.docx'
        format_document(draft,out,m)
        d=Document(out);self.assertEqual(audit(d,m,out)['errors'],0)
        self.assertTrue(any(p.text.startswith('    return') for p in paragraphs(d)))
    def test_no_unscoped_normalization(self):
        src,work,mf=self.make_template()
        with self.assertRaises(ValueError):format_document(work,self.root/'bad.docx')
    def test_outline_parent_and_question(self):
        entries=[{'index':0,'level':1,'question':'做什么','parent_index':None},
                 {'index':1,'level':3,'question':'如何实现','parent_index':0}]
        self.assertTrue(validate_outline(entries))
        entries[1]['level']=2
        self.assertFalse(validate_outline(entries))
        entries[1]['question']=''
        self.assertTrue(validate_outline(entries))
    def test_body_number_is_not_heading(self):
        d=Document();ensure_styles(d)
        p=d.add_paragraph('1. 这是步骤，不是章节',style='SZUBody')
        self.assertIsNone(heading_level(p))
    def test_mutated_heading_properties(self):
        d=Document();ensure_styles(d);n=new_numbering(d)
        p=d.add_paragraph('实验方法');apply_heading(p,1,n)
        p=d.add_paragraph('具体方法');apply_heading(p,2,n)
        p.runs[0].bold=True;p.paragraph_format.left_indent=Pt(30)
        rules={i['rule'] for i in audit(d)['issues']}
        self.assertIn('heading_bold',rules);self.assertIn('heading_indent',rules)
    def test_evidence_and_review_never_auto_pass(self):
        f=self.root/'evidence.json';write_json(f,{'requirements':[{'id':'R1','status':'complete','needs_result':True,'evidence':[]}]})
        self.assertTrue(evidence_audit(f))
        self.assertFalse(validate_review(None,'unused',{})[0])
    def test_border_override_detected(self):
        d=Document();ensure_styles(d);t=d.add_table(rows=2,cols=2);format_table(t,'data')
        node(node(t.cell(0,0)._tc.get_or_add_tcPr(),'tcBorders'),'bottom',val='nil')
        self.assertIn('cell_border_override',{i['rule'] for i in audit(d)['issues']})
    def test_rendered_numbering_catches_word_only_failure(self):
        render={'backend':'word','headings':[{'level':1,'label':'一、','text':'方法','page':1},{'level':2,'label':'','text':'实现','page':1}]}
        self.assertTrue(rendered_numbering_issues(render))
        render['headings'][1]['label']='(一)'
        self.assertFalse(rendered_numbering_issues(render))
    def test_explicit_outline_rewrites_wrong_existing_heading(self):
        d=Document();d.add_paragraph('9. 原先错误的标题');src=self.root/'draft.docx';d.save(src)
        plan={'paragraphs':[{'index':0,'text':'9. 原先错误的标题','replacement_text':'实验方法','level':1,'parent_index':None,'question':'采用什么方法'}]}
        out=self.root/'restructured.docx';format_document(src,out,outline_data=plan,generated_all=True)
        result=Document(out);self.assertEqual(result.paragraphs[0].text,'实验方法');self.assertEqual(audit(result)['errors'],0)
    def test_semantic_parent_context_is_preserved(self):
        d=Document();t=d.add_table(rows=1,cols=1);t.cell(0,0).text='实验过程'
        src=self.root/'parent.docx';d.save(src);work=self.root/'work.docx';mf=self.root/'m.json'
        prepare(src,work,mf,layout_tables=[0],heading_parents=['0:1'])
        m=load_manifest(mf);d=Document(work);ensure_styles(d);p=d.tables[0].cell(0,0).add_paragraph('方法实现');apply_heading(p,2,new_numbering(d))
        self.assertEqual(audit(d,m)['errors'],0)
    def test_teacher_region_rejects_new_signature(self):
        d=Document();t=d.add_table(rows=2,cols=1);t.cell(0,0).text='实验过程';t.cell(1,0).text='指导教师签字：'
        src=self.root/'t.docx';d.save(src);work=self.root/'w.docx';mf=self.root/'m.json';prepare(src,work,mf,layout_tables=[0])
        d=Document(work);d.tables[0].cell(1,0).add_paragraph('代填签名')
        self.assertIn('generated_content_in_teacher_region',{i['rule'] for i in audit_protection(d,load_manifest(mf))})
    def test_template_defaults_are_protected(self):
        src,work,mf=self.make_template();m=load_manifest(mf);d=Document(work)
        default=d.styles.element.find(qn('w:docDefaults'));node(node(default,'rPrDefault'),'rPr').append(OxmlElement('w:b'))
        self.assertIn('protected_document_defaults_changed',{i['rule'] for i in audit_protection(d,m)})


def visual_fixture(output):
    """Long prose, multilevel restarts, a wrapping heading, code, figures, long table."""
    from PIL import Image,ImageDraw
    out=Path(output);out.parent.mkdir(parents=True,exist_ok=True)
    d=Document();s=d.sections[0]
    from docx.shared import Mm
    s.page_width=Mm(210);s.page_height=Mm(297);s.left_margin=s.right_margin=Mm(25.4)
    ensure_styles(d);n=new_numbering(d);page_frame(d)
    def heading(title,level):
        p=d.add_paragraph(title);apply_heading(p,level,n)
    def body(t):
        p=d.add_paragraph(style='SZUBody');set_font(p.add_run(t),10.5,latin='Times New Roman')
    heading('实验方法与实现',1)
    heading('标题分级与内容组织',2)
    heading('依据正文包含关系建立标题层级并检查多行标题在实际打印页面中的缩进和换行对齐效果',3)
    heading('结构检查方法',4)
    heading('检查父子关系',5);body('本节验证标题对应的主题、层级与段落内容一致。脚本检查格式，语义仍须逐节审查。')
    heading('检查同级关系',5);body('同级标题采用相同的划分依据，避免将整体流程与某个参数并列。')
    heading('视觉检查方法',4)
    heading('检查段落缩进',5);body('本节验证上级变化后五级编号重新开始，正文首行保持两个字符缩进。')
    add_code_box(d,'def evaluate(result):\n    return result is not None')
    body('evaluate接收结果对象，返回结果是否存在。此代码仅为排版测试，不是用户实验结果。')
    heading('长正文与分页',2)
    for i in range(8):body(('分页测试用于确认长正文能自然延续，标题不会单独留在页底。外框应闭合且不截断内容。'*4))
    heading('实验结果与分析',1)
    heading('数据表格检查',2)
    t=d.add_table(rows=36,cols=3);format_table(t,'data')
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            p=c.paragraphs[0];p.style='SZUTable'
            p.text=(['编号','测试项目','结果'][j] if i==0 else [str(i),'跨页表格的表头重复与网格连续性','通过'][j])
            for r in p.runs:set_font(r,10.5,latin='Times New Roman')
    p=d.add_paragraph('表1分页验证数据',style='SZUCaption')
    heading('图形与题注检查',2)
    img=Image.new('RGB',(650,220),'white');draw=ImageDraw.Draw(img)
    draw.rectangle((20,20,630,200),outline='black',width=3)
    for i in range(8):draw.rectangle((40+i*70,180-i*15,80+i*70,185),fill='#333333')
    image_path=out.with_suffix('.png');img.save(image_path)
    p=d.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(image_path),width=Mm(130))
    d.add_paragraph('图1排版测试示意图',style='SZUCaption')
    body('此图仅用于验证图片与题注同页以及图中文字和线条清晰度。')
    heading('结论',1);body('样稿用于回归验证，不对应用户实验内容。只有逐页检查完成后，才能确认视觉要求满足。')
    d.save(out)


if __name__=='__main__':
    import sys
    if len(sys.argv)>1 and sys.argv[1]=='--visual-fixture':visual_fixture(sys.argv[2])
    else:unittest.main()
