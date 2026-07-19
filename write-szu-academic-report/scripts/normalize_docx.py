#!/usr/bin/env python3
"""Normalize common SZU report formatting while preserving document structure."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

CAPTION_RE = re.compile(r"^\s*图\s*\d+")
HEADING_RE = re.compile(r"^(?:第[一二三四五六七八九十]+[章节]|[一二三四五六七八九十]+、|\d+(?:\.\d+){0,3})\s*[^。；，]{1,40}$")


def style_name(paragraph) -> str:
    return paragraph.style.name if paragraph.style is not None else ""


def heading_level(paragraph) -> int | None:
    name = style_name(paragraph)
    for level in range(1, 5):
        if str(level) in name and ("Heading" in name or "标题" in name):
            return level
    m = re.match(r"^(\d+(?:\.\d+){0,3})", paragraph.text.strip())
    if m and HEADING_RE.match(paragraph.text.strip()):
        return min(4, m.group(1).count(".") + 1)
    if re.match(r"^(?:第[一二三四五六七八九十]+[章节]|[一二三四五六七八九十]+、)", paragraph.text.strip()):
        return 1
    return None


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def set_run(run, size: float | None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = None


def set_paragraph_runs(paragraph, size: float | None) -> None:
    for run in paragraph.runs:
        set_run(run, size)


def normalize(path: Path, output: Path) -> None:
    doc = Document(path)
    seen_heading = False
    sizes = {1: 14.0, 2: 12.0, 3: 10.5, 4: 9.0}

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        level = heading_level(paragraph)
        if has_drawing(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if CAPTION_RE.match(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
            set_paragraph_runs(paragraph, 10.5)
            continue
        if level is not None:
            seen_heading = True
            paragraph.paragraph_format.first_line_indent = None
            set_paragraph_runs(paragraph, sizes[level])
            continue
        if seen_heading and text:
            paragraph.paragraph_format.first_line_indent = Pt(21)
            set_paragraph_runs(paragraph, 10.5)
        else:
            set_paragraph_runs(paragraph, None)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = None
                    set_paragraph_runs(paragraph, 10.5)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    normalize(args.input, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
